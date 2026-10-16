"""Excel export for a selected optimization run."""
from io import BytesIO
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def schedule_xlsx(report, experiment, result):
    if not result.get('assignment'):
        raise ValueError('No schedule is available for this run.')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Optimized schedule'
    ws.append(['exam_id', 'student_count', 'reference_start', 'optimized_start', 'optimized_end'])
    for exam, slot in zip(report['exams'], result['assignment']):
        start = datetime.fromisoformat(report['starts'][slot])
        ws.append([exam['exam_id'], exam['student_count'], report['starts'][exam['reference_slot']],
                   start.isoformat(timespec='minutes'), (start + timedelta(minutes=report['dataset']['duration_minutes'])).isoformat(timespec='minutes')])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='126B70')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in 'ACDE':
        ws.column_dimensions[col].width = 25
    ws.column_dimensions['B'].width = 18
    meta = wb.create_sheet('Experiment')
    for row in [('Application', 'ExamFlow'), ('Source', report['dataset']['source_name']),
                ('Scenario', experiment['scenario']), ('Capacity', experiment['capacity']),
                ('Method', result['algorithm']), ('Seed', result['seed']),
                ('Duration (min)', report['dataset']['duration_minutes']), ('Peak', result['metrics']['peak'])]:
        meta.append(row)
    meta.column_dimensions['A'].width = 22
    meta.column_dimensions['B'].width = 48
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
