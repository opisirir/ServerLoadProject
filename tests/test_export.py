import json
import subprocess
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from examlab.export import schedule_xlsx


class ExportTests(unittest.TestCase):
    def test_python_and_browser_exports_preserve_schedule(self):
        report = json.loads(Path('web/data/benchmark.json').read_text())
        e = report['experiments'][0]
        run = next(r for r in e['runs'] if r.get('assignment'))
        with tempfile.TemporaryDirectory() as d:
            target = Path(d) / 'browser.xlsx'
            subprocess.run(['node', '-e', "const fs=require('fs'),x=require('./web/xlsx.js'),r=JSON.parse(fs.readFileSync('web/data/benchmark.json'));const e=r.experiments[0],s=e.runs.find(r=>r.assignment);fs.writeFileSync(process.argv[1],x.schedule(r,e,s));", str(target)], check=True)
            for data in [schedule_xlsx(report,e,run),target.read_bytes()]:
                wb = load_workbook(BytesIO(data))
                ws = wb['Optimized schedule']
                self.assertEqual(ws.max_row,len(report['exams'])+1)
                self.assertEqual(ws['D2'].value,report['starts'][run['assignment'][0]])
                self.assertEqual(sum(row[1] for row in ws.iter_rows(min_row=2,values_only=True)), report['dataset']['enrollments'])
                self.assertEqual(wb['Experiment']['B4'].value,e['capacity'])
                self.assertEqual(ws.freeze_panes,'A2')
